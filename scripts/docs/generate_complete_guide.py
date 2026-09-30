# ============================================================================
# SCRIPT PARA GENERAR GUÍA EXPLICATIVA COMPLETA EN WORD (.DOCX)
# Proyecto: Reconocimiento de Emociones Faciales con Deep Learning (Emotion AI v2.0)
# Autores: Manuel Fabrizzio Risco Gil & Alejandro Choquihuayta, Mario Julian
# Curso: Sistemas Inteligentes y Machine Learning (Ciclo 9)
# ============================================================================

import os
import sys

# Asegurar que el directorio raíz del proyecto esté en sys.path
BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, BASE_DIR)

from src.utils import setup_utf8
setup_utf8()

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

OUTPUT_DIR = os.path.join(BASE_DIR, "outputs")
DOC_PATH = os.path.join(OUTPUT_DIR, "Guia_Explicativa_Completa_Proyecto_Emociones_CNN.docx")
SEMANA1_DIR = os.path.dirname(BASE_DIR)
DOC_PATH_SEMANA1 = os.path.join(SEMANA1_DIR, "Guia_Explicativa_Completa_Proyecto_Emociones_CNN.docx")


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def add_heading_styled(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = "Calibri"
        if level == 1:
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 51, 102)
        elif level == 2:
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 80, 160)
        elif level == 3:
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 120, 215)
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(4)
    return heading


def add_body_text(doc, text, bold=False, italic=False, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(40, 40, 40)
    p.alignment = alignment
    pf = p.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15
    if alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
        pf.first_line_indent = Cm(0.75)
    return p


def add_bullet(doc, title, desc):
    p = doc.add_paragraph(style='List Bullet')
    r1 = p.add_run(title + ": ")
    r1.font.name = "Calibri"
    r1.font.size = Pt(10.5)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0, 51, 102)
    r2 = p.add_run(desc)
    r2.font.name = "Calibri"
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = RGBColor(40, 40, 40)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    return p


def add_callout_box(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "EBF3FA")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_title = p.add_run("💡 " + title + "\n")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(11)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0, 51, 102)

    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor(30, 30, 30)
    p.paragraph_format.space_after = Pt(2)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(4)
    p_space.paragraph_format.space_after = Pt(4)


def format_table_header(row, col_names, bg_color="003366"):
    for idx, name in enumerate(col_names):
        cell = row.cells[idx]
        cell.text = name
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)


def format_table_row(row, values, is_even=False, align_center_cols=None):
    bg_color = "F4F7FA" if is_even else "FFFFFF"
    align_center_cols = align_center_cols or []
    for idx, val in enumerate(values):
        cell = row.cells[idx]
        cell.text = str(val)
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in align_center_cols else WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(40, 40, 40)


def generate_comprehensive_guide():
    print("\n" + "=" * 60)
    print("   GENERANDO GUÍA EXPLICATIVA EN WORD (.DOCX)...")
    print("=" * 60)

    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # =========================================================================
    # PORTADA / ENCABEZADO
    # =========================================================================
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(25)
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run("GUÍA EXPLICATIVA COMPLETA - EMOTION AI v2.0")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0, 51, 102)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(18)
    r_sub = p_sub.add_run("Reconocimiento de Emociones Faciales con ResNet-SE, YuNet ONNX, VRAM y Distribución Portátil")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(12.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(70, 70, 70)

    table_meta = doc.add_table(rows=5, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_info = [
        ("Autores / Estudiantes:", "• Manuel Fabrizzio Risco Gil\n• Alejandro Choquihuayta, Mario Julian"),
        ("Curso y Ciclo:", "Sistemas Inteligentes y Machine Learning (9no Ciclo)"),
        ("Arquitectura Principal:", "EmotionResNetSE (Residual Blocks + Squeeze-and-Excitation Attention)"),
        ("Detector Facial:", "YuNet Deep Learning Face Detector (cv2.FaceDetectorYN - Multi-Ángulo)"),
        ("Despliegue Portátil:", "Exportación ONNX universal + OpenCV DNN (Lanzador con 1 clic en cualquier PC)")
    ]
    for idx, (label, val) in enumerate(meta_info):
        row = table_meta.rows[idx]
        c1, c2 = row.cells[0], row.cells[1]
        c1.text = label
        c2.text = val
        set_cell_background(c1, "EAEFF5")
        set_cell_background(c2, "F8FAFC")
        set_cell_margins(c1, 60, 60, 100, 100)
        set_cell_margins(c2, 60, 60, 100, 100)
        c1.paragraphs[0].runs[0].font.bold = True
        c1.paragraphs[0].runs[0].font.size = Pt(10)
        c2.paragraphs[0].runs[0].font.size = Pt(10)

    doc.add_page_break()

    # =========================================================================
    # SECCIÓN 1: INTRODUCCIÓN Y CONCEPTOS BÁSICOS
    # =========================================================================
    add_heading_styled(doc, "1. ¿Qué es este proyecto y qué problemas resuelve?", level=1)

    add_body_text(doc,
        "Este proyecto implementa un sistema inteligente de visión artificial capaz de clasificar 7 emociones humanas "
        "(Enojado, Disgusto, Miedo, Feliz, Triste, Sorprendido y Neutral) en imágenes estáticas y video en tiempo real. "
        "La versión v2.0 resuelve los desafíos críticos de los sistemas tradicionales mediante innovaciones clave:"
    )

    add_bullet(doc, "1. Detección Facial Neuronal Multi-Ángulo", "Utiliza YuNet ONNX, detectando rostros en perfiles de hasta ±90° de giro, roll y pitch, con 5 puntos faciales (landmarks) y sin falsos positivos en ropa o fondos.")
    add_bullet(doc, "2. Clasificación de Expresiones Sutiles", "Introduce EmotionResNetSE con atención por canales (Squeeze-and-Excitation) para priorizar cejas, ojos y comisuras labiales.")
    add_bullet(doc, "3. Manejo del Desbalance Severo", "Combina Focal Loss (gamma=2.0) con Oversampling estocástico, elevando el Recall en clases raras como 'Disgusto' a más del 72.9%.")
    add_bullet(doc, "4. Inferencia Fluida sin Parpadeos (Jitter)", "Aplica Suavizado Temporal EMA (Exponential Moving Average) para generar transiciones de probabilidad suaves y estables en tiempo real.")
    add_bullet(doc, "5. Despliegue Universal Portátil", "Exporta el modelo a ONNX y compila un paquete portátil autónomo (.ZIP) que corre en cualquier PC con puro OpenCV y NumPy a 60 FPS.")

    add_callout_box(doc, "Idea Clave: ¿Cómo funciona el aprendizaje con 35K imágenes?",
        "Mediante Data Augmentation acelerado en GPU (Rotación Afín ±12°, Random Crop con padding reflejado, Flips y Cutout) "
        "y regularización Mixup (alpha=0.1), la red visualiza millones de variantes únicas en memoria a lo largo del entrenamiento, "
        "generalizando de forma altamente robusta."
    )

    # =========================================================================
    # SECCIÓN 2: ARQUITECTURA EMOTIONRESNETSE
    # =========================================================================
    add_heading_styled(doc, "2. Arquitectura de la Red Neuronal: EmotionResNetSE", level=1)

    add_body_text(doc,
        "EmotionResNetSE cuenta con 11,389,383 parámetros organizados en 4 etapas residuales con módulos de atención por canales:"
    )

    table_arch = doc.add_table(rows=8, cols=4)
    table_arch.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_table_header(table_arch.rows[0], ["Capa / Módulo", "Operaciones Internas", "Dimensión Salida", "Parámetros"])

    arch_data = [
        ("Stem Inicial", "Conv2D(1->64, 3x3) + BatchNorm + ReLU", "64 × 48 × 48", "704"),
        ("Etapa 1 (Stage 1)", "2x ResidualSEBlock(64->64) + SE-Attention", "64 × 48 × 48", "150,016"),
        ("Etapa 2 (Stage 2)", "2x ResidualSEBlock(64->128, stride=2) + SE-Attention", "128 × 24 × 24", "529,664"),
        ("Etapa 3 (Stage 3)", "2x ResidualSEBlock(128->256, stride=2) + SE-Attention", "256 × 12 × 12", "2,116,096"),
        ("Etapa 4 (Stage 4)", "2x ResidualSEBlock(256->512, stride=2) + SE-Attention", "512 × 6 × 6", "8,459,264"),
        ("Global Avg Pool", "AdaptiveAvgPool2D(1, 1) + Flatten", "Vector (512)", "0"),
        ("Clasificador FC", "Dropout(0.4) + Linear(512->256) + BN1D + ReLU + Dropout(0.3) + Linear(256->7)", "Vector (7)", "133,639"),
    ]
    for i, row_data in enumerate(arch_data):
        format_table_row(table_arch.rows[i+1], row_data, is_even=(i % 2 == 1), align_center_cols=[0, 2, 3])

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(4)

    add_heading_styled(doc, "¿Por qué el mecanismo SE-Attention marca la diferencia?", level=2)
    add_body_text(doc,
        "El módulo Squeeze-and-Excitation (SE Block) analiza qué canales de características contienen información relevante "
        "para la emoción actual. Por ejemplo, al sonreír, los canales que detectan la curvatura de la boca reciben una mayor "
        "ponderación multiplicativa (peso cercano a 1.0), mientras que canales con ruido de fondo se atenúan (peso cercano a 0.0)."
    )

    # =========================================================================
    # SECCIÓN 3: PROCESAMIENTO 100% EN VRAM Y FOCAL LOSS
    # =========================================================================
    add_heading_styled(doc, "3. Procesamiento en VRAM, Focal Loss y Balanceo", level=1)

    add_body_text(doc,
        "El sistema aprovecha al máximo la tarjeta gráfica NVIDIA GeForce RTX 5060 Ti mediante tres pilares:"
    )
    add_bullet(doc, "Dataset 100% Residente en VRAM", "Las 35,887 imágenes se transfieren como tensores uint8 a la memoria GPU (78.9 MB totales), eliminando latencias del bus PCIe y permitiendo más de 12,000 imgs/seg.")
    add_bullet(doc, "Focal Loss (gamma=2.0)", "Modula la función de pérdida para concentrar el gradiente en las emociones difíciles y suprimir la sobreconfianza en las fáciles, resolviendo el ratio 17:1 entre Feliz y Disgusto.")
    add_bullet(doc, "Planificador SequentialLR", "8 épocas de Linear Warmup para estabilizar pesos iniciales seguidas de decaimiento Cosine Annealing para alcanzar mínimos globales estables.")

    # =========================================================================
    # SECCIÓN 4: RESULTADOS VERIFICADOS EN TEST SET
    # =========================================================================
    add_heading_styled(doc, "4. Resultados Experimentales en el Conjunto de TEST (7,178 Muestras)", level=1)

    add_body_text(doc,
        "El modelo final EmotionResNetSE fue evaluado formalmente sobre las 7,178 imágenes independientes del conjunto de Test, "
        "obteniendo una Exactitud Global (Accuracy) de 69.45% (4,985 aciertos de 7,178) y superando el promedio de concordancia humana (65%-68%):"
    )

    table_test = doc.add_table(rows=10, cols=5)
    table_test.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_table_header(table_test.rows[0], ["Emoción", "Precision", "Recall", "F1-Score", "Muestras"])

    test_res = [
        ("Enojado", "61.55%", "62.84%", "62.19%", "958"),
        ("Disgusto", "61.65%", "73.87%", "67.21%", "111"),
        ("Miedo", "57.92%", "49.61%", "53.45%", "1,024"),
        ("Feliz", "87.92%", "88.61%", "88.27%", "1,774"),
        ("Triste", "61.21%", "50.12%", "55.11%", "1,247"),
        ("Sorprendido", "77.25%", "84.60%", "80.76%", "831"),
        ("Neutral", "60.71%", "72.42%", "66.05%", "1,233"),
        ("PROMEDIO MACRO", "66.89%", "68.87%", "67.58%", "7,178"),
        ("PROMEDIO PONDERADO", "69.17%", "69.45%", "69.05%", "7,178")
    ]
    for i, row_data in enumerate(test_res):
        format_table_row(table_test.rows[i+1], row_data, is_even=(i % 2 == 1), align_center_cols=[1, 2, 3, 4])

    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(4)

    # Insertar gráficos clave en la guía
    plots_to_insert = [
        ("historial_entrenamiento.png", "Figura 1: Curvas de Aprendizaje (Pérdida y Exactitud) a lo largo de 96 épocas.", 6.0),
        ("matriz_confusion_normalizada.png", "Figura 2: Matriz de Confusión Normalizada (%) sobre el conjunto de Test.", 5.0),
        ("f1_precision_recall_por_emocion.png", "Figura 3: Comparativa de F1-Score y Precisión vs. Recall por emoción.", 5.8)
    ]
    plots_dir = os.path.join(OUTPUT_DIR, "plots")
    for fname, cap, w_in in plots_to_insert:
        p_path = os.path.join(plots_dir, fname)
        if os.path.exists(p_path):
            try:
                p_im = doc.add_paragraph()
                p_im.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_im.paragraph_format.space_before = Pt(8)
                p_im.paragraph_format.space_after = Pt(2)
                p_im.add_run().add_picture(p_path, width=Inches(w_in))
                p_cp = doc.add_paragraph()
                p_cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_c = p_cp.add_run(cap)
                r_c.font.name = "Calibri"
                r_c.font.size = Pt(9.5)
                r_c.font.italic = True
                p_cp.paragraph_format.space_after = Pt(8)
            except Exception:
                pass

    # =========================================================================
    # SECCIÓN 5: SISTEMA EN TIEMPO REAL (WEBCAM v2.0)
    # =========================================================================
    add_heading_styled(doc, "5. Sistema en Tiempo Real: YuNet, EMA Tracker y 60 FPS", level=1)

    add_body_text(doc,
        "El módulo src/realtime.py ofrece una experiencia de producción optimizada:"
    )
    add_bullet(doc, "Detector Facial YuNet ONNX", "Detección profunda multi-ángulo (yaw ±90°, inclinación y cabeceo) con umbral de confianza (>0.65) y 5 puntos faciales (landmarks).")
    add_bullet(doc, "Suavizado Temporal (EMA Tracker)", "Fórmula exponencial P_t = 0.30 * P_actual + 0.70 * P_anterior. Evita parpadeos bruscos por microexpresiones.")
    add_bullet(doc, "Preprocesamiento C++ Ultrarrápido", "Redimensiona el recorte del rostro en C++ directamente a tensores en <0.05 ms.")
    add_bullet(doc, "Optimización de Hardware MJPG", "Configura la cámara web a 60 FPS y buffersize=1, eliminando retrasos por cola de fotogramas.")
    add_bullet(doc, "Ventana Dinámica Ajustable", "Ventana libremente redimensionable mediante mouse o maximizado con cv2.WINDOW_NORMAL.")

    # Captura en tiempo real
    sc_path = os.path.join(plots_dir, "screenshot_516.png")
    if os.path.exists(sc_path):
        try:
            p_im = doc.add_paragraph()
            p_im.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_im.paragraph_format.space_before = Pt(8)
            p_im.paragraph_format.space_after = Pt(2)
            p_im.add_run().add_picture(sc_path, width=Inches(5.0))
            p_cp = doc.add_paragraph()
            p_cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_c = p_cp.add_run("Figura 4: Interfaz de inferencia en tiempo real con HUD holográfico y detector YuNet.")
            r_c.font.name = "Calibri"
            r_c.font.size = Pt(9.5)
            r_c.font.italic = True
            p_cp.paragraph_format.space_after = Pt(8)
        except Exception:
            pass

    # =========================================================================
    # SECCIÓN 6: EXPORTACIÓN PORTÁTIL (.ZIP PARA CUALQUIER PC)
    # =========================================================================
    add_heading_styled(doc, "6. Exportación Portátil Universal (Opción 5 del Menú)", level=1)

    add_body_text(doc,
        "Para compartir el proyecto y que cualquier persona o jurado pueda probarlo en su propia computadora "
        "sin necesidad de instalar PyTorch ni contar con tarjeta gráfica dedicada, se diseñó la Opción 5 del menú:"
    )

    add_bullet(doc, "¿Cómo se comparte?", "Se copia el archivo 'Emotion_AI_Portable_v2.0.zip' (40.5 MB) a un USB o Google Drive.")
    add_bullet(doc, "¿Cómo se ejecuta en otra PC?", "El usuario solo debe descomprimir el ZIP y hacer doble clic en 'EJECUTAR_WEBCAM.bat'.")
    add_bullet(doc, "¿Qué dependencias requiere?", "Únicamente Python con OpenCV y NumPy (se autoinstalan en segundos). No requiere los 3 GB de PyTorch CUDA.")
    add_bullet(doc, "¿Cómo corre la red neuronal?", "OpenCV DNN carga directamente el modelo 'modelo_emociones_cv2.onnx' y ejecuta la inferencia en la CPU a más de 30-60 FPS.")

    # =========================================================================
    # SECCIÓN 7: PREGUNTAS CLAVE PARA SUSTENTACIÓN
    # =========================================================================
    add_heading_styled(doc, "7. Preguntas Clave para Sustentación o Examen", level=1)

    qa_list = [
        ("¿Por qué el modelo alcanza 69.45% en Test y supera el benchmark humano?",
         "Porque incorpora conexiones residuales (ResNet) para evitar el desvanecimiento de gradiente, módulos de atención SE para priorizar rasgos morfológicos clave, y Focal Loss con Oversampling para equilibrar emociones difíciles."),
        ("¿Cómo se comprobó que no hay fuga de datos (Data Leakage)?",
         "Las 7,178 imágenes de Test se aislaron completamente y nunca se usaron en entrenamiento. El Data Augmentation y Oversampling solo se aplican sobre el conjunto de entrenamiento (is_train=True)."),
        ("¿Qué ventajas ofrece YuNet frente a clasificadores tradicionales?",
         "YuNet es una red convolucional profunda ligera (ONNX) que detecta caras de perfil (yaw ±90°), cabeceo e inclinación, extrae 5 puntos faciales y elimina falsos positivos en ropa y fondos."),
        ("¿En qué consiste el Suavizado Temporal EMA en tiempo real?",
         "Es una media móvil exponencial sobre las probabilidades de los rostros rastreados por IoU. Filtra el ruido inter-frame y erradica el parpadeo errático de las barras de probabilidad."),
        ("¿Cómo se logró que el programa funcione en cualquier computadora sin tarjeta gráfica NVIDIA?",
         "Se exportó el modelo PyTorch al estándar universal ONNX. Esto permite ejecutar la red mediante el motor OpenCV DNN nativo en CPU con solo 15 MB de librerías, empaquetado en un archivo .ZIP con un ejecutable .bat de 1 solo clic.")
    ]

    for q, a in qa_list:
        add_bullet(doc, "P: " + q, "\nR: " + a)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    doc.save(DOC_PATH)
    try:
        doc.save(DOC_PATH_SEMANA1)
    except Exception:
        pass

    print(f"\n[OK] Guía explicativa generada exitosamente en:")
    print(f"     1. {DOC_PATH}")
    print(f"     2. {DOC_PATH_SEMANA1}")


if __name__ == "__main__":
    generate_comprehensive_guide()
